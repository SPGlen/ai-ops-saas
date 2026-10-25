"""Admission templates and resumable, reviewed table imports. No scanned mark books money."""

import io, json, zipfile, hashlib, subprocess, sys, tempfile, threading, os
from pathlib import Path
from datetime import datetime
from sqlmodel import Session, select
from sqlalchemy import update
from .models import *
from .db import engine, set_scope
from .config import ROOT, settings
from .domain import query, own, plain_data, create_record, update_record
from .permissions import Scope, permit, identity_for
from .security import fail
from .studio import seal, unseal, field_access, values_for, definition_version
from .operations import event
from .integrations.storage import Storage, save_asset, upload_mime

COLUMNS = [
    ("external_row_id", "线索编号"),
    ("student_name", "学员"),
    ("grade", "年级"),
    ("name", "家长"),
    ("relationship", "关系"),
    ("contact", "电话"),
    ("source", "来源"),
    ("class_name", "意向班级"),
    ("owner_name", "招生负责人"),
    ("trial_at", "预约时间"),
    ("stage", "跟进阶段"),
    ("registration_note", "报名信息"),
    ("loss_reason", "流失原因"),
    ("notes", "备注"),
]
STAGES = {
    "新线索": "new",
    "已联系": "contacted",
    "已预约": "booked",
    "已到场": "attended",
    "已报名": "won",
    "暂不报名": "lost",
    "流失": "lost",
    "退费": "withdrawn",
}
LOCK = threading.Lock()


def template_fields(i, definition_id=""):
    fields = [{"key": k, "label": label, "type": "text"} for k, label in COLUMNS]
    version = "admissions-1"
    if definition_id:
        d = own(i.session, StudioDefinition, i.tenant.id, definition_id)
        if d.kind != "form" or d.binding != "lead":
            fail("SCHEMA_INVALID", "请选择已发布的招生表单")
        v = definition_version(i.session, i.tenant.id, d)
        version = d.id + ":" + str(v.version)
        fields += [
            f
            for f in v.spec.get("fields", [])
            if f["key"] not in {x["key"] for x in fields} and field_access(f, i)
        ]
    return fields, version


def xlsx_template(i, definition_id=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    fields, version = template_fields(i, definition_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "招生录入"
    count = len(fields)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=count)
    ws.cell(1, 1, i.tenant.name + " · 招生信息登记表")
    ws.cell(1, 1).font = Font(name="微软雅黑", size=18, bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="194E3C")
    ws.row_dimensions[1].height = 38
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=count)
    ws.cell(
        2,
        1,
        "模板版本："
        + version
        + "    金额及收退费状态仅供核对，正式状态以系统记录为准。",
    )
    ws.row_dimensions[2].height = 26
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=count)
    ws.cell(
        3,
        1,
        "请保留线索编号和表头。扫描识别后核对确认；重复导入不会自动覆盖更新后的记录。",
    )
    ws.row_dimensions[3].height = 26
    prefix = "R" + uid()[:8]
    for col, f in enumerate(fields, 1):
        c = ws.cell(4, col, f["label"])
        c.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="31614B")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        from openpyxl.utils import get_column_letter

        ws.column_dimensions[get_column_letter(col)].width = (
            20 if f["key"] in ("trial_at", "loss_reason", "notes") else 17
        )
        for row in range(5, 25):
            c = ws.cell(
                row, col, (prefix + "-" + str(row - 4).zfill(3)) if col == 1 else None
            )
            c.font = Font(
                name="微软雅黑", size=11, color="2563A5" if col > 1 else "526B60"
            )
            c.fill = PatternFill(
                "solid", fgColor="F0F5EF" if row % 2 == 0 else "FFFFFF"
            )
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = Border(bottom=Side(style="thin", color="D9E4DC"))
            c.number_format = "@"
            ws.row_dimensions[row].height = 30
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{ws.cell(24, count).column_letter}24"
    ws.sheet_view.showGridLines = False
    dv = DataValidation(type="list", formula1='"新线索,已联系,已预约,已到场,暂不报名"')
    ws.add_data_validation(dv)
    stagecol = next(j + 1 for j, f in enumerate(fields) if f["key"] == "stage")
    dv.add(f"{ws.cell(5, stagecol).coordinate}:{ws.cell(24, stagecol).coordinate}")
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def pdf_template(i, definition_id=""):
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib import colors
    import qrcode
    from reportlab.lib.utils import ImageReader

    fields, version = template_fields(i, definition_id)
    if len(fields) > 22:
        fail("PRINT_TOO_WIDE", "打印模板最多22列，请减少显示字段或使用Excel")
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    out = io.BytesIO()
    w, h = landscape(A3)
    c = Canvas(out, pagesize=(w, h))
    c.setTitle("招生信息登记表")
    c.setFillColor(colors.HexColor("#194E3C"))
    c.rect(28, h - 82, w - 56, 52, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("STSong-Light", 22)
    c.drawString(44, h - 62, i.tenant.name + " · 招生信息登记表")
    c.setFillColor(colors.HexColor("#415D50"))
    c.setFont("STSong-Light", 9)
    c.drawString(
        28,
        h - 101,
        "模板版本："
        + version
        + "  |  请保留编号；扫描后核对，收退费状态以系统记录为准。",
    )
    image = io.BytesIO()
    qrcode.make("zhiying-admissions:" + version).save(image, format="PNG")
    image.seek(0)
    c.drawImage(ImageReader(image), w - 81, h - 79, 46, 46)
    weights = [
        1.5
        if f["key"]
        in (
            "contact",
            "trial_at",
            "notes",
            "loss_reason",
            "registration_note",
            "external_row_id",
        )
        else 1
        for f in fields
    ]
    widths = [(w - 56) * v / sum(weights) for v in weights]
    top = h - 121
    rh = 37
    prefix = "R" + uid()[:8]
    for row in range(17):
        x = 28
        y = top - row * rh
        for idx, f in enumerate(fields):
            c.setStrokeColor(colors.HexColor("#BDCFC1"))
            c.setFillColor(colors.HexColor("#E5EFE2") if row == 0 else colors.white)
            c.rect(x, y - rh, widths[idx], rh, fill=1, stroke=1)
            c.setFillColor(colors.HexColor("#183E2D"))
            c.setFont("STSong-Light", 9 if row == 0 else 8)
            text = (
                f["label"]
                if row == 0
                else (prefix + "-" + str(row).zfill(3) if idx == 0 else "")
            )
            c.drawString(x + 4, y - 23, text)
            x += widths[idx]
    c.setFont("STSong-Light", 9)
    c.drawString(
        28,
        32,
        "填写日期：____________    录入人：____________    编号用于匹配原记录，请勿重复用于新的家庭。",
    )
    c.save()
    return out.getvalue()


def read_excel(raw):
    from openpyxl import load_workbook

    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(x.file_size for x in archive.infolist()) > 80 * 1024 * 1024 or any(
            "vbaproject" in x.filename.lower() for x in archive.infolist()
        ):
            fail("FILE_INVALID", "Excel内容过大或包含宏")
    wb = load_workbook(
        io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
    )
    ws = wb.worksheets[0]
    rows = []
    headers = None
    version = "admissions-1"
    aliases = {label: key for key, label in COLUMNS}
    aliases.update(
        {
            "联系方式": "contact",
            "家长称呼": "name",
            "学员姓名": "student_name",
            "孩子姓名": "student_name",
        }
    )
    for cells in ws.iter_rows():
        values = [str(c.value).strip() if c.value is not None else "" for c in cells]
        for v in values:
            if "模板版本：" in v:
                version = v.split("模板版本：", 1)[1].split()[0]
        if headers is None:
            if ("学员" in values or "学员姓名" in values) and (
                "电话" in values or "联系方式" in values
            ):
                headers = {
                    j: aliases.get(v, "custom:" + v) for j, v in enumerate(values) if v
                }
            continue
        data = {headers[j]: v for j, v in enumerate(values) if j in headers and v}
        if not any(data.get(k) for k in ("student_name", "name", "contact")):
            continue
        issues = (
            ["包含公式，请手动填写实际值"]
            if any(c.data_type == "f" for c in cells)
            else []
        )
        rows.append(
            {
                "data": data,
                "issues": issues,
                "confidence": 1.0,
                "page": 1,
                "source_row": cells[0].row,
            }
        )
        if len(rows) > 500:
            fail("ROWS_EXCEEDED", "最多导入500行")
    wb.close()
    if headers is None:
        fail("TEMPLATE_INVALID", "未识别到招生表头，请使用系统模板")
    return rows, version


def read_scan(raw, name, i=None, template_version="admissions-1"):
    from pypdf import PdfReader

    if raw.startswith(b"%PDF"):
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted or len(reader.pages) > 20:
            fail("PAGES_EXCEEDED", "PDF最多20页，且不能加密")
    pages = recognize(raw, name)
    rows = []
    aliases = {label: key for key, label in COLUMNS}
    for page in pages:
        for word in page:
            if word["text"].startswith("__TEMPLATE_VERSION__:"):
                template_version = word["text"].split(":", 1)[1]
    if i and template_version != "admissions-1":
        try:
            fid, number = template_version.rsplit(":", 1)
            definition = own(i.session, StudioDefinition, i.tenant.id, fid)
            version = definition_version(
                i.session, i.tenant.id, definition, int(number)
            )
        except (ValueError, AttributeError):
            fail("TEMPLATE_INVALID", "模板版本无法核对，请重新下载")
        if definition.kind != "form" or definition.binding != "lead":
            fail("TEMPLATE_INVALID", "请选择本机构招生表单")
        aliases.update(
            {
                f["label"].replace(" ", ""): "custom:" + f["label"]
                for f in version.spec.get("fields", [])
                if f["key"] not in dict(COLUMNS)
            }
        )
    for pi, words in enumerate(pages, 1):
        headers = []
        for word in words:
            label = word["text"].strip().replace(" ", "")
            if label in aliases:
                headers.append({**word, "key": aliases[label]})
        if not {"student_name", "contact"} <= set(h["key"] for h in headers):
            continue
        headers.sort(key=lambda w: w["x"])
        hy = max(h["y"] for h in headers)
        values = [w for w in words if w["y"] > hy + 10]
        rowgroups = []
        for word in sorted(values, key=lambda w: w["y"]):
            if not rowgroups or abs(word["y"] - rowgroups[-1]["y"]) > max(
                12, word.get("height", 12) * 0.8
            ):
                rowgroups.append({"y": word["y"], "words": []})
            rowgroups[-1]["words"].append(word)
        for group in rowgroups:
            d = {}
            scores = []
            cell_confidence = {}
            for word in sorted(group["words"], key=lambda w: w["x"]):
                col = min(headers, key=lambda h: abs(h["x"] - word["x"]))["key"]
                d[col] = (d.get(col, "") + " " + word["text"]).strip()
                scores.append(word["score"])
                cell_confidence[col] = min(cell_confidence.get(col, 1), word["score"])
            if not any(d.get(k) for k in ("student_name", "name", "contact")):
                continue
            confidence = min(scores or [0])
            rows.append(
                {
                    "data": d,
                    "confidence": confidence,
                    "cell_confidence": cell_confidence,
                    "issues": ["扫描结果需要逐项核对"]
                    + (["识别置信度较低"] if confidence < 0.9 else []),
                    "page": pi,
                    "source_row": len(rows) + 1,
                }
            )
            if len(rows) > 500:
                fail("ROWS_EXCEEDED", "最多500行")
    if not rows:
        fail("TEMPLATE_INVALID", "未识别到表格行，请拍摄清晰原图或使用Excel导入")
    return rows, template_version


def preview(i, rows):
    scope = Scope(i)
    visible = {r.id: r for r in scope.list("lead")}
    allleads = [r for r in scope.rows.values() if r.kind == "lead"]
    for row in rows:
        d = row["data"]
        row["cell_issues"] = {
            key: "识别置信度低，请核对"
            for key, score in row.get("cell_confidence", {}).items()
            if score < 0.9
        }
        key = d.get("external_row_id", "")
        matched = next(
            (
                r
                for r in allleads
                if r.id == key or (key and r.data.get("external_row_id") == key)
            ),
            None,
        )
        row.update(match_id="", base_version=None, status="pending", candidates=[])
        if matched and matched.id not in visible:
            row["issues"].append("该编号已有记录，需负责人处理")
            row["blocked"] = True
        elif matched:
            row.update(match_id=matched.id, base_version=matched.version)
            if d.get("student_name") and d["student_name"] != matched.data.get(
                "student_name"
            ):
                row["issues"].append("编号相同但学员姓名改变，请核对是否为新家庭")
        else:
            for r in visible.values():
                plain = plain_data(i.tenant.id, r.data)
                if d.get("contact") and plain.get("contact") == d["contact"]:
                    row["candidates"].append(
                        {
                            "id": r.id,
                            "name": plain.get("student_name", plain["name"]),
                            "version": r.version,
                        }
                    )
            if row["candidates"]:
                row["issues"].append(
                    "发现相同家长电话的记录，请核对；不同孩子不会自动合并"
                )
        if d.get("stage") in ("已报名", "退费", "won", "withdrawn") or d.get(
            "registration_note"
        ):
            row["issues"].append("报名/退费信息仅作为核对资料，需在系统完成业务确认")
    return rows


def execute(batch_id, t):
    with Session(engine(), expire_on_commit=False) as s:
        set_scope(s, t)
        batch = s.exec(
            query(s, AdmissionImport, t)
            .where(AdmissionImport.id == batch_id)
            .with_for_update()
        ).one()
        if batch.status != "queued":
            return
        claim = s.exec(
            update(AdmissionImport)
            .where(
                AdmissionImport.id == batch.id,
                AdmissionImport.tenant_id == t,
                AdmissionImport.status == "queued",
            )
            .values(status="processing", progress=10, updated_at=now())
        )
        s.commit()
        if claim.rowcount != 1:
            return
        s.refresh(batch)
        try:
            tenant = s.get(Tenant, t)
            i = identity_for(s, tenant, batch.actor_id)
            permit(i, "lead.write")
            from .auth import writable

            writable(i)
            asset = own(s, Asset, t, batch.asset_id)
            raw = Storage().read(asset.object_key)
            rows, version = (
                read_excel(raw)
                if asset.mime.endswith("spreadsheetml.sheet")
                else read_scan(raw, asset.name, i, batch.template_version)
            )
            s.expire_all()
            tenant = s.get(Tenant, t)
            i = identity_for(s, tenant, batch.actor_id)
            permit(i, "lead.write")
            writable(i)
            batch.rows_cipher = seal(t, preview(i, rows))
            batch.template_version = version
            batch.updated_at = now()
            batch.status = "review"
            batch.progress = 100
            batch.version += 1
            s.add(batch)
            event(s, t, batch.actor_id, "import.review", batch.id)
            s.commit()
        except Exception as exc:
            s.rollback()
            batch = s.get(AdmissionImport, batch_id)
            detail = getattr(exc, "detail", {})
            batch.status = "failed"
            batch.error = (
                detail.get("message", "识别未完成，请检查文件和OCR配置")
                if isinstance(detail, dict)
                else "识别未完成"
            )
            s.add(batch)
            s.commit()


def confirm(i, batch, edits, version):
    if batch.actor_id != i.user.id and "owner" not in i.roles:
        fail("NOT_FOUND", "导入记录不属于您", 404)
    permit(i, "lead.write")
    if version != batch.version:
        fail("VERSION_CONFLICT", "导入结果已更新，请刷新", 409)
    if batch.status not in ("review", "partial", "completed"):
        fail("STATE_INVALID", "请先等待识别完成")
    rows = unseal(i.tenant.id, batch.rows_cipher)
    results = []
    for edit in edits:
        index = edit["index"]
        if not 0 <= index < len(rows):
            fail("ROW_INVALID", "行号错误")
        row = rows[index]
        if row.get("status") == "imported":
            results.append(
                {
                    "index": index,
                    "status": "already_imported",
                    "id": row.get("record_id"),
                }
            )
            continue
        try:
            with i.session.begin_nested():
                d = dict(edit.get("data", row["data"]))
                matched = edit.get("match_id", row.get("match_id", ""))
                if row.get("blocked"):
                    fail("ROLE_FORBIDDEN", "需负责人核对编号", 403)
                old = Scope(i).own(matched, "lead", "write") if matched else None
                if old and old.version != (
                    row.get("base_version")
                    if old.id == row.get("match_id")
                    else next(
                        (
                            c["version"]
                            for c in row.get("candidates", [])
                            if c["id"] == old.id
                        ),
                        -1,
                    )
                ):
                    fail("VERSION_CONFLICT", "原线索已更新，请重新核对", 409)
                key = d.get("external_row_id", "")
                if key and not old:
                    duplicate = next(
                        (
                            r
                            for r in i.session.exec(
                                query(i.session, Record, i.tenant.id).where(
                                    Record.kind == "lead"
                                )
                            ).all()
                            if r.data.get("external_row_id") == key
                        ),
                        None,
                    )
                    if duplicate:
                        fail("DUPLICATE_ROW", "该编号已经导入，请刷新核对", 409)
                stage = STAGES.get(d.get("stage", ""), d.get("stage", "new"))
                if stage in ("won", "withdrawn"):
                    stage = old.data.get("stage", "new") if old else "new"
                values = {
                    k: d.get(k, "")
                    for k in (
                        "student_name",
                        "name",
                        "relationship",
                        "contact",
                        "grade",
                        "source",
                        "loss_reason",
                        "notes",
                    )
                }
                values["stage"] = stage
                if (
                    not values["name"]
                    or not values["contact"]
                    or not values["student_name"]
                ):
                    fail("REQUIRED", "请填写学员、家长和电话")
                if old:
                    values = {**plain_data(i.tenant.id, old.data), **values}
                    values = {
                        k: v
                        for k, v in values.items()
                        if k
                        in __import__(
                            "app.schemas", fromlist=["Lead"]
                        ).Lead.model_fields
                    }
                if d.get("registration_note"):
                    values["notes"] = (
                        values.get("notes", "")
                        + "\n纸表报名信息（待核对）："
                        + d["registration_note"]
                    )[:2000]
                if d.get("class_name"):
                    cls = next(
                        (
                            r
                            for r in Scope(i).rows.values()
                            if r.kind == "class" and r.data["name"] == d["class_name"]
                        ),
                        None,
                    )
                    if not cls:
                        fail("REFERENCE_INVALID", "意向班级名称未匹配")
                    values["class_id"] = cls.id
                if d.get("owner_name") and "owner" in i.roles:
                    from .permissions import member_roles

                    members = i.session.exec(
                        select(Membership).where(Membership.tenant_id == i.tenant.id)
                    ).all()
                    users = [
                        i.session.get(User, m.user_id)
                        for m in members
                        if member_roles(m) & {"owner", "admissions"}
                    ]
                    users = [u for u in users if u and u.name == d["owner_name"]]
                    if len(users) != 1:
                        fail("REFERENCE_INVALID", "招生负责人名称需要唯一匹配")
                    values["owner_id"] = users[0].id
                if d.get("trial_at"):
                    try:
                        values["trial_at"] = datetime.fromisoformat(
                            d["trial_at"]
                        ).isoformat()
                    except ValueError:
                        fail("DATE_INVALID", "预约日期格式错误")
                values = Scope(i).validate_write("lead", values, old)
                from .studio import guard_action

                guard_action(i, "lead", old.id if old else "", values, "")
                record = (
                    update_record(
                        i.session, i.tenant.id, i.user.id, old.id, values, old.version
                    )
                    if old
                    else create_record(
                        i.session, i.tenant.id, i.user.id, "lead", values
                    )
                )
                record.data = {
                    **record.data,
                    "external_row_id": key,
                    "import_batch_id": batch.id,
                }
                i.session.add(record)
                i.session.flush()
                custom = {
                    str(k)[7:]: v
                    for k, v in d.items()
                    if str(k).startswith("custom:") and v
                }
                if custom:
                    try:
                        definition_id, number = batch.template_version.rsplit(":", 1)
                        definition = own(
                            i.session, StudioDefinition, i.tenant.id, definition_id
                        )
                        form_version = definition_version(
                            i.session, i.tenant.id, definition, int(number)
                        )
                    except (ValueError, AttributeError):
                        fail("TEMPLATE_INVALID", "自定义列需要使用已发布表单模板")
                    if definition.kind != "form" or definition.binding != "lead":
                        fail("TEMPLATE_INVALID", "自定义模板不属于招生业务")
                    fields = form_version.spec.get("fields", [])
                    labels = {f["label"]: f for f in fields}
                    if set(custom) - set(labels):
                        fail("FORM_INVALID", "存在未定义的自定义列")
                    typed = {}
                    for label, value in custom.items():
                        field = labels[label]
                        if field["type"] == "multiselect" and isinstance(value, str):
                            import re

                            value = [
                                x.strip()
                                for x in re.split(r"[,，;；\n]", value)
                                if x.strip()
                            ]
                        if field["type"] == "boolean" and isinstance(value, str):
                            if value.strip().lower() in (
                                "是",
                                "true",
                                "1",
                                "已完成",
                                "√",
                            ):
                                value = True
                            elif value.strip().lower() in (
                                "否",
                                "false",
                                "0",
                                "未完成",
                            ):
                                value = False
                        typed[field["key"]] = value
                    values = values_for(i, fields, typed, True)
                    submission = FormSubmission(
                        tenant_id=i.tenant.id,
                        definition_id=definition.id,
                        definition_version=form_version.version,
                        actor_id=i.user.id,
                        values_cipher=seal(i.tenant.id, values),
                        binding="lead",
                        record_id=record.id,
                        request_id="import:" + batch.id + ":" + str(index),
                    )
                    i.session.add(submission)
                    i.session.flush()
                    if form_version.spec.get("workflow_id"):
                        from .studio import submit

                        flow = own(
                            i.session,
                            StudioDefinition,
                            i.tenant.id,
                            form_version.spec["workflow_id"],
                        )
                        inst = submit(
                            i,
                            flow,
                            {"_form_submission_id": submission.id, "values": values},
                            "",
                            "form:" + submission.id,
                            form_version.spec.get("workflow_version"),
                        )
                        submission.status = inst.status
                        i.session.add(submission)
                row.update(status="imported", record_id=record.id, data=d)
                results.append({"index": index, "status": "imported", "id": record.id})
        except Exception as exc:
            detail = getattr(exc, "detail", {})
            message = (
                detail.get("message", "该行未完成，请检查数据")
                if isinstance(detail, dict)
                else "该行未完成"
            )
            row.update(status="conflict", error=message)
            results.append({"index": index, "status": "conflict", "error": message})
    batch.rows_cipher = seal(i.tenant.id, rows)
    batch.status = (
        "completed" if all(r.get("status") == "imported" for r in rows) else "partial"
    )
    batch.version += 1
    i.session.add(batch)
    event(
        i.session,
        i.tenant.id,
        i.user.id,
        "import.confirmed",
        batch.id,
        {"imported": sum(r.get("status") == "imported" for r in rows)},
    )
    return {"results": results, "version": batch.version, "status": batch.status}


def recognize(raw, name):
    with tempfile.TemporaryDirectory(prefix="ops-table-") as directory:
        source = Path(directory) / (
            "scan.pdf" if raw.startswith(b"%PDF") else "scan.png"
        )
        source.write_bytes(raw)
        interpreter = settings().ocr_python or sys.executable
        env = {
            k: v
            for k, v in os.environ.items()
            if k.upper()
            in (
                "PATH",
                "SYSTEMROOT",
                "WINDIR",
                "TEMP",
                "TMP",
                "USERPROFILE",
                "APPDATA",
                "LOCALAPPDATA",
                "HOME",
                "SSL_CERT_FILE",
                "SSL_CERT_DIR",
            )
        }
        env["PYTHONIOENCODING"] = "utf-8"
        env["MODELSCOPE_CACHE"] = str(settings().data_dir / "ocr-cache/modelscope")
        env["HF_HOME"] = str(settings().data_dir / "ocr-cache/huggingface")
        env["PADDLE_PDX_CACHE_HOME"] = str(settings().data_dir / "ocr-cache")
        env["PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK"] = "True"
        with LOCK:
            result = subprocess.run(
                [
                    interpreter,
                    str(ROOT / "backend/app/table_ocr_runner.py"),
                    str(source),
                ],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=600,
                env=env,
            )
        marker = "OPS_OCR_RESULT="
        if result.returncode or marker not in result.stdout:
            fail(
                "OCR_UNAVAILABLE",
                "结构化OCR未完成，请检查专用OCR运行环境或改用Excel导入",
                503,
            )
        return json.loads(result.stdout.rsplit(marker, 1)[1])
