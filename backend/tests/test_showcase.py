import base64
import pytest
from sqlmodel import Session, select
from app.config import settings
from app.db import engine
from app.models import Tenant, LedgerEntry, CreditEntry, NotificationOutbox, PublishJob, Record, now
from .conftest import login, add


def demo_env(monkeypatch, code=""):
    monkeypatch.setenv("OPS_DEMO_MODE", "true")
    monkeypatch.setenv("OPS_TASK_MODE", "demo")
    monkeypatch.setenv("OPS_DEMO_EXPIRES_AT", "")
    monkeypatch.setenv("OPS_DEMO_ACCESS_CODE", code)
    monkeypatch.setenv("OPS_SMTP_HOST", "")
    monkeypatch.setenv("OPS_S3_BUCKET", "")
    monkeypatch.setenv("OPS_PUBLISHER_URL", "")
    monkeypatch.setenv("OPS_ALLOW_LIVE_PUBLISH", "false")
    monkeypatch.setenv("OPS_ALLOW_LIVE_REPLIES", "false")
    settings.cache_clear()


def test_showcase_is_opt_in(client):
    assert client.get("/api/v1/demo/info").json()=={"enabled":False}
    assert client.get("/api/v1/demo/showcase").status_code==404
    assert client.get("/api/v1/auth/me").json()["features"]["demo_mode"] is False


def test_demo_external_gate_and_normal_business_auth(client,monkeypatch):
    demo_env(monkeypatch,"a-private-showcase-access-2026")
    response=client.get("/")
    assert response.status_code==200 and 'id="demo-access-form"' in response.text
    assert "www-authenticate" not in response.headers
    response=client.get("/api/v1/demo/info")
    assert response.status_code==401 and "www-authenticate" not in response.headers
    # A previously cached native-browser login must not trigger any challenge.
    client.headers["Authorization"]="Basic "+base64.b64encode("演示:incorrect".encode()).decode()
    assert client.get("/").status_code==200
    origin={"Origin":settings().public_url}
    response=client.post("/__demo/access",json={"code":"incorrect"},headers=origin)
    assert response.status_code==401 and "www-authenticate" not in response.headers
    response=client.post("/__demo/access",json={"code":"a-private-showcase-access-2026"},headers=origin)
    assert response.status_code==200,response.text
    assert "HttpOnly" in response.headers["set-cookie"] and "SameSite=strict" in response.headers["set-cookie"]
    assert client.get("/api/v1/demo/info").status_code==200
    assert 'id="demo-access-form"' not in client.get("/").text
    assert "noindex" in client.get("/").headers["x-robots-tag"]
    for method,path in [("put","/integrations/aliyun"),("post","/runs"),("post","/automation-runs"),("post","/notifications/test"),("post","/social-accounts/a/login"),("post","/auth/register")]:
        response=getattr(client,method)("/api/v1"+path,json={})
        assert response.status_code==409,response.text
        assert response.json()["detail"]["code"]=="DEMO_EXTERNAL_DISABLED"
    lead=add(client,"lead",{"name":"演示线索","contact":"DEMO-ONLY"})
    assert lead["data"]["name"]=="演示线索"
    client.cookies.delete("ops_session")
    assert client.get("/api/v1/demo/showcase").json()["detail"]["code"]=="AUTH_REQUIRED"
    client.cookies.clear()
    assert client.get("/api/v1/demo/showcase").json()["detail"]["code"]=="DEMO_ACCESS_REQUIRED"


def test_demo_web_gate_rejects_wrong_origin_tampering_and_rotated_code(client,monkeypatch):
    from app.demo_access import token, COOKIE
    demo_env(monkeypatch,"a-private-showcase-access-2026")
    cfg=settings()
    response=client.post("/__demo/access",json={"code":cfg.demo_access_code},headers={"Origin":"https://foreign.example"})
    assert response.status_code==403 and COOKIE not in client.cookies
    response=client.post("/__demo/access",content='x'*1100,headers={"Origin":cfg.public_url,"Content-Type":"application/json"})
    assert response.status_code==413
    client.cookies.set(COOKIE,token(cfg)+"tampered")
    assert client.get("/api/v1/demo/info").status_code==401
    client.cookies.set(COOKIE,token(cfg))
    assert client.get("/api/v1/demo/info").status_code==200
    monkeypatch.setenv("OPS_DEMO_ACCESS_CODE","rotated-private-showcase-access-2026")
    settings.cache_clear()
    assert client.get("/api/v1/demo/info").status_code==401
    # No role records or credentials are exposed by the public readiness route.
    assert client.get("/api/v1/health").status_code==200


def test_demo_web_gate_rate_limits_wrong_codes_and_cookie_expires(client,monkeypatch):
    from datetime import timedelta
    from app import demo_access
    demo_env(monkeypatch,"a-private-showcase-access-2026")
    cfg=settings();origin={"Origin":cfg.public_url}
    for n in range(20):
        assert client.post("/__demo/access",json={"code":"wrong"},headers=origin).status_code==401
    response=client.post("/__demo/access",json={"code":"wrong"},headers=origin)
    assert response.status_code==429 and "www-authenticate" not in response.headers
    old=now()-timedelta(hours=9)
    with monkeypatch.context() as clock:
        clock.setattr(demo_access,"now",lambda:old)
        value=demo_access.token(cfg)
    client.cookies.set(demo_access.COOKIE,value)
    assert client.get("/api/v1/demo/info").status_code==401


def test_demo_requires_separate_database_and_public_access_code(client,monkeypatch):
    from app.demo_runtime import check_demo_database
    demo_env(monkeypatch)
    with pytest.raises(RuntimeError,match="non-showcase"):
        check_demo_database()
    monkeypatch.setenv("OPS_PUBLIC_URL","https://demo.example")
    settings.cache_clear()
    with pytest.raises(RuntimeError,match="access code"):
        settings().validate_runtime()
    monkeypatch.setenv("OPS_DEMO_ACCESS_CODE","a-private-showcase-access-2026")
    settings.cache_clear();settings().validate_runtime()
    monkeypatch.setenv("OPS_ALIYUN_API_KEY","a-real-key-must-not-be-accepted")
    settings.cache_clear()
    with pytest.raises(RuntimeError,match="real model"):
        settings().validate_runtime()


def test_demo_overview_preserves_tenant_boundary(client,monkeypatch):
    demo_env(monkeypatch)
    add(client,"lead",{"name":"机构A线索","contact":"DEMO-A"})
    first=client.get("/api/v1/demo/showcase").json()
    assert first["counts"]["lead"]==1
    login(client,"b")
    second=client.get("/api/v1/demo/showcase").json()
    assert second["counts"]["lead"]==0
    assert second["external_calls"] is False


def test_fixture_generation_is_replayable_and_journals_are_real(tmp_path,monkeypatch):
    monkeypatch.setenv("OPS_DATABASE_URL","sqlite:///"+str(tmp_path/"showcase.db").replace("\\","/"))
    monkeypatch.setenv("OPS_DATA_DIR",str(tmp_path/"data"))
    monkeypatch.setenv("OPS_SECRET_KEY","demo-fixture-unique-test-key-at-least-32-chars")
    monkeypatch.setenv("OPS_ENVIRONMENT","development")
    monkeypatch.setenv("OPS_ALIYUN_API_KEY","")
    monkeypatch.setenv("OPS_PUBLIC_URL","http://127.0.0.1:8188")
    demo_env(monkeypatch)
    engine.cache_clear()
    from app import demo_showcase
    from app.permissions import identity_for
    from app.models import User, now
    from app.reports import build, TZ
    monkeypatch.setattr(demo_showcase,"preview_video",lambda *args:None)
    demo_showcase.seed()
    with Session(engine()) as s:
        t=s.exec(select(Tenant).where(Tenant.slug=="qinghe-demo")).one()
        u=s.exec(select(User).where(User.email=="owner.qinghe@zhiying.example")).one()
        report=build(identity_for(s,t,u.id),"monthly",now().astimezone(TZ).strftime("%Y-%m"))
        assert report["metrics"]["population"]==30
        assert report["metrics"]["attrition_rate"]==30
        assert report["metrics"]["actual_attendance"]>0
        assert report["metrics"]["refund_count"]==9
        assert report["finance"]["totals"]["refund_fen"]==360000
        counts=(len(s.exec(select(LedgerEntry)).all()),len(s.exec(select(CreditEntry)).all()))
        from app.domain import plain_data
        from app.models import StudioDefinition, FormSubmission, FlowTask
        for record in s.exec(select(Record).where(Record.kind.in_(["student", "lead"]))).all():
            data=plain_data(record.tenant_id, record.data)
            assert len(data["contact"]) == 11 and data["contact"].isdigit()
            assert "contact" not in record.data  # persisted encrypted, including demo values
            for guardian in data.get("guardians", []):
                assert len(guardian["contact"]) == 11 and guardian["contact"].isdigit()
        assert len(s.exec(select(Record).where(Record.tenant_id==t.id, Record.kind=="community")).all())==3
        assert len(s.exec(select(Record).where(Record.tenant_id==t.id, Record.kind=="topic")).all())==6
        assert len(s.exec(select(FormSubmission).where(FormSubmission.tenant_id==t.id)).all())==2
        assert any(task.status=="pending" for task in s.exec(select(FlowTask).where(FlowTask.tenant_id==t.id)).all())
        assert not any(d.enabled for d in s.exec(select(StudioDefinition).where(StudioDefinition.kind=="workflow")).all())
        assert len(s.exec(select(PublishJob)).all())==0
        assert all(n.status=="disabled" for n in s.exec(select(NotificationOutbox)).all())
    demo_showcase.seed()
    with Session(engine()) as s:
        assert counts==(len(s.exec(select(LedgerEntry)).all()),len(s.exec(select(CreditEntry)).all()))
    engine().dispose();engine.cache_clear();settings.cache_clear()


def test_showcase_switching_has_separate_rate_limit(client,monkeypatch):
    from app import routes_auth
    calls=[]
    monkeypatch.setattr(routes_auth,"limit",lambda session,key,maximum,seconds:calls.append((key,maximum,seconds)))
    login(client)
    assert calls[-1][1:]==(20,900)
    demo_env(monkeypatch)
    login(client)
    assert calls[-1][1:]==(60,60)
    assert calls[-1][0].startswith("showcase-login:")
    # Default workspace connections must not leak into a showcase launch,
    # including Windows PowerShell 5 which removes blank env overrides.
    monkeypatch.delenv("OPS_PUBLISHER_URL", raising=False)
    settings.cache_clear()
    assert not settings().publisher_url



def test_demo_expired_link_is_closed(client,monkeypatch):
    demo_env(monkeypatch)
    monkeypatch.setenv("OPS_DEMO_EXPIRES_AT","2020-01-01T00:00:00+00:00")
    settings.cache_clear()
    assert client.get("/").status_code==410
    assert client.get("/api/v1/demo/info").status_code==410


def test_demo_excel_import_works_and_scan_requires_real_ocr(client,monkeypatch):
    from app import routes_operations, admission_import
    from app.models import AdmissionImport
    from io import BytesIO
    from openpyxl import Workbook
    demo_env(monkeypatch)
    wb=Workbook();ws=wb.active
    # Use the exact product template so version and header validation remain active.
    response=client.get("/api/v1/admissions/template/xlsx")
    assert response.status_code==200
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(response.content));ws=wb.active
    for col,(key,label) in enumerate(admission_import.COLUMNS,1):
        values={"external_row_id":"DEMO-IMPORT-001","student_name":"演示体验学员","name":"演示家长","contact":"DEMO-0001","grade":"二年级","source":"企业体验","stage":"已联系"}
        ws.cell(5,col,values.get(key,""))
    out=BytesIO();wb.save(out)
    queued=[];monkeypatch.setattr(routes_operations,"dispatch_import",lambda tid,bid:queued.append((tid,bid)))
    response=client.post("/api/v1/admissions/imports",files={"file":("demo.xlsx",out.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code==201,response.text
    assert queued
    admission_import.execute(queued[0][1],queued[0][0])
    with Session(engine()) as s:
        batch=s.get(AdmissionImport,queued[0][1]);assert batch.status=="review",batch.error
    from PIL import Image
    out=BytesIO();Image.new("RGB",(16,16)).save(out,format="PNG")
    response=client.post("/api/v1/admissions/imports",files={"file":("scan.png",out.getvalue(),"image/png")})
    assert response.status_code==409
    assert response.json()["detail"]["code"]=="DEMO_OCR_PAUSED"
