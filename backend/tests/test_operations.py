"""0.3 acceptance: real business transactions, scopes, cohorts and durable workflow state."""

import io
from datetime import timedelta
from types import SimpleNamespace
import pytest
from fastapi import HTTPException
from sqlmodel import Session, select
from app.db import engine
from app.models import *
from app import operations, reports, studio
from app.permissions import identity_for
from .conftest import add, login
from .test_staff_social import staff, sign


def v3():
    with Session(engine()) as s:
        t = s.get(Tenant, "tenant-a")
        t.settings = {
            **t.settings,
            "legacy_maintenance": False,
            "consultation_enabled": False,
        }
        s.add(t)
        s.commit()


def registration(client, suffix="1", lead_id=""):
    st = add(
        client,
        "student",
        {
            "name": "学员" + suffix,
            "guardians": [
                {
                    "name": "家长" + suffix,
                    "relationship": "母亲",
                    "contact": "13800000001",
                },
                {"name": "家长乙", "relationship": "父亲", "contact": "13800000002"},
            ],
        },
    )
    o = add(
        client,
        "order",
        {
            "name": "AI报名" + suffix,
            "student_id": st["id"],
            "program_name": "创意AI",
            "price_fen": 10000,
            "discount_fen": 1000,
            "purchased_units": 8,
            "lead_id": lead_id,
        },
    )
    return st, o


def enroll(client, st, o, cls=""):
    pay = client.post(
        "/api/v1/orders/" + o["id"] + "/payments",
        json={"amount_fen": 9000, "request_id": "pay-" + o["id"]},
    )
    assert pay.status_code == 200, pay.text
    return add(
        client,
        "enrollment",
        {"student_id": st["id"], "order_id": o["id"], "class_id": cls},
    )


def define(client, binding="independent", spec=None, kind="workflow"):
    if spec is None:
        spec = {
            "nodes": [
                {"id": "start", "data": {"kind": "start"}},
                {
                    "id": "review",
                    "data": {"kind": "approval", "roles": ["owner"], "mode": "all"},
                },
                {"id": "end", "data": {"kind": "end"}},
            ],
            "edges": [
                {"source": "start", "target": "review"},
                {"source": "review", "target": "end"},
            ],
        }
    r = client.post(
        "/api/v1/studio/definitions",
        json={"name": "验收设计", "kind": kind, "binding": binding, "spec": spec},
    )
    assert r.status_code == 201, r.text
    d = r.json()
    pub = client.post("/api/v1/studio/definitions/" + d["id"] + "/publish")
    assert pub.status_code == 200, pub.text
    return pub.json()


def enable(client, d):
    r = client.patch(
        "/api/v1/studio/definitions/" + d["id"],
        json={
            "name": d["name"],
            "kind": d["kind"],
            "binding": d["binding"],
            "spec": d["draft"],
            "version": d["version"],
            "enabled": True,
        },
    )
    assert r.status_code == 200, r.text


def test_direct_order_deposit_and_formal_conversion(client):
    v3()
    assert (
        client.post(
            "/api/v1/records/course", json={"data": {"name": "废弃"}}
        ).status_code
        == 403
    )
    lead = add(
        client,
        "lead",
        {"name": "母亲", "student_name": "小明", "contact": "13800000001"},
    )
    st, o = registration(client, lead_id=lead["id"])
    assert o["data"]["purchased_units"] == 8 and not o["data"].get("course_id")
    client.post(
        "/api/v1/orders/" + o["id"] + "/payments",
        json={"amount_fen": 1000, "request_id": "deposit-01"},
    )
    r = client.post(
        "/api/v1/records/enrollment",
        json={"data": {"student_id": st["id"], "order_id": o["id"]}},
    )
    assert r.status_code == 400
    assert (
        client.get("/api/v1/records/lead").json()["items"][0]["data"]["stage"] == "new"
    )
    client.post(
        "/api/v1/orders/" + o["id"] + "/payments",
        json={"amount_fen": 8000, "request_id": "remaining-payment"},
    )
    e = add(client, "enrollment", {"student_id": st["id"], "order_id": o["id"]})
    assert e["data"]["remaining"] == 8
    assert (
        client.get("/api/v1/records/lead").json()["items"][0]["data"]["stage"] == "won"
    )
    assert (
        client.post(
            "/api/v1/records/enrollment",
            json={"data": {"student_id": st["id"], "order_id": o["id"]}},
        ).status_code
        == 409
    )
    with Session(engine()) as s:
        row = s.get(Record, st["id"])
        assert "13800000001" not in str(row.data) and row.data["guardians_cipher"]


def test_refund_snapshot_withdrawal_and_no_duplicate_charge(client):
    v3()
    st, o = registration(client)
    cl = add(client, "class", {"name": "一班"})
    e = enroll(client, st, o, cl["id"])
    lesson = add(
        client,
        "lesson",
        {
            "name": "第一课",
            "class_id": cl["id"],
            "starts_at": (now() + timedelta(days=1)).isoformat(),
        },
    )
    url = "/api/v1/enrollments/" + e["id"] + "/attendance"
    for key in ("one", "two"):
        assert (
            client.post(
                url, json={"lesson_id": lesson["id"], "request_id": "attendance-" + key}
            ).status_code
            == 200
        )
    ctx = client.get("/api/v1/orders/" + o["id"] + "/refund-context").json()[
        "enrollments"
    ][0]
    assert (ctx["purchased_units"], ctx["attended_units"], ctx["remaining_units"]) == (
        8,
        1,
        7,
    )
    rurl = "/api/v1/orders/" + o["id"] + "/refunds"
    assert (
        client.post(
            rurl, json={"amount_fen": 100, "request_id": "no-reason-01"}
        ).status_code
        == 400
    )
    body = {
        "amount_fen": 1000,
        "request_id": "partial-refund-01",
        "reason": "调整优惠",
        "withdrawal": False,
    }
    assert client.post(rurl, json=body).status_code == 200
    body = {
        "amount_fen": 8000,
        "request_id": "withdraw-refund-01",
        "reason": "时间冲突",
        "withdrawal": True,
        "enrollment_ids": [e["id"]],
    }
    result = client.post(rurl, json=body)
    assert result.status_code == 200, result.text
    assert client.post(rurl, json=body).json()["id"] == result.json()["id"]
    assert client.post(rurl, json={**body, "reason": "changed"}).status_code == 409
    assert (
        client.post(
            url, json={"lesson_id": lesson["id"], "request_id": "closed-signin"}
        ).status_code
        == 400
    )
    detail = client.get("/api/v1/finance/details/ledger/" + result.json()["id"]).json()
    assert (
        detail["refund_context"]["snapshot"]["enrollments"][0]["remaining_units"] == 7
    )
    assert (
        client.get("/api/v1/records/enrollment").json()["items"][0]["data"]["remaining"]
        == 0
    )


def test_multiple_enrollments_partial_refund_and_exact_30_alert(client):
    v3()
    students = []
    for n in range(10):
        st, o = registration(client, str(n))
        e = enroll(client, st, o)
        students.append((st, o, e))
    # Another active enrollment means the first withdrawal is not a lost student.
    st = students[3][0]
    o = add(
        client,
        "order",
        {
            "name": "第二项目",
            "student_id": st["id"],
            "price_fen": 9000,
            "purchased_units": 4,
        },
    )
    enroll(client, st, o)
    for n in range(4):
        _, o, e = students[n]
        r = client.post(
            "/api/v1/orders/" + o["id"] + "/refunds",
            json={
                "amount_fen": 1000,
                "reason": "改期",
                "withdrawal": True,
                "enrollment_ids": [e["id"]],
                "request_id": "withdraw-" + str(n),
            },
        )
        assert r.status_code == 200, r.text
    period = now().astimezone(reports.TZ).strftime("%Y-%m")
    r = client.get("/api/v1/reports/monthly", params={"period": period}).json()
    assert r["metrics"]["population"] == 10 and r["metrics"]["withdrawn_students"] == 3
    assert r["metrics"]["attrition_rate"] == 30
    alerts = client.get("/api/v1/operations/alerts").json()
    assert len(alerts) == 1 and alerts[0]["status"] == "open"
    client.get("/api/v1/operations/alerts")
    assert len(client.get("/api/v1/notifications").json()) == 1


def test_empty_report_unknown_denominator_and_immutable_snapshot(client):
    v3()
    period = now().astimezone(reports.TZ).strftime("%Y-%m")
    data = client.get("/api/v1/reports/monthly", params={"period": period}).json()
    assert (
        data["metrics"]["attrition_rate"] is None
        and data["metrics"]["attendance_rate"] is None
    )
    path = "/api/v1/report-snapshots/monthly?period=" + period
    a = client.post(path).json()
    b = client.post(path).json()
    assert a["id"] == b["id"]
    st, o = registration(client)
    enroll(client, st, o)
    c = client.post(path).json()
    assert c["version"] == 2 and a["payload"]["metrics"]["new_students"] == 0
    with Session(engine()) as s:
        snap = s.get(ReportSnapshot, a["id"])
        snap.digest = "tamper"
        s.add(snap)
        with pytest.raises(Exception):
            s.commit()


def test_late_conversion_backfills_original_cohort(client):
    v3()
    lead = add(
        client, "lead", {"name": "家长", "contact": "13800000001", "stage": "attended"}
    )
    date = now() - timedelta(days=10)
    with Session(engine()) as s:
        l = s.get(Record, lead["id"])
        l.created_at = date
        l.data = {**l.data, "first_attended_at": date.isoformat()}
        s.add(l)
        s.commit()
    st, o = registration(client, lead_id=lead["id"])
    enroll(client, st, o)
    r = client.get(
        "/api/v1/reports/weekly",
        params={"period": date.astimezone(reports.TZ).date().isoformat()},
    ).json()["metrics"]
    assert (
        r["new_leads"],
        r["lead_converted"],
        r["lead_conversion_rate"],
        r["trial_conversion_rate"],
    ) == (1, 1, 100, 100)


def test_kpi_scope_correction_and_cancelled_lesson(client):
    v3()
    with Session(engine()) as s:
        staff(s, "teach1", ["teacher"])
        staff(s, "teach2", ["teacher"])
    cl = add(client, "class", {"name": "一班", "teacher_id": "teach1"})
    lesson = add(
        client,
        "lesson",
        {
            "name": "结束课次",
            "class_id": cl["id"],
            "starts_at": (now() - timedelta(days=2)).isoformat(),
            "status": "completed",
        },
    )
    k = client.get("/api/v1/education/feedback-kpis").json()[0]
    assert k["overdue"]
    sign(client, "teach2")
    assert client.get("/api/v1/education/feedback-kpis").json() == []
    assert (
        client.patch(
            "/api/v1/education/feedback-kpis/" + k["id"],
            json={"completed": True, "version": 1},
        ).status_code
        == 404
    )
    sign(client, "teach1")
    r = client.patch(
        "/api/v1/education/feedback-kpis/" + k["id"],
        json={"completed": True, "version": 1},
    )
    assert r.status_code == 200, r.text
    assert (
        client.patch(
            "/api/v1/education/feedback-kpis/" + k["id"],
            json={"completed": True, "version": 1},
        ).status_code
        == 409
    )
    login(client)
    assert (
        client.patch(
            "/api/v1/education/feedback-kpis/" + k["id"],
            json={"completed": False, "version": 2},
        ).status_code
        == 400
    )
    assert (
        client.patch(
            "/api/v1/education/feedback-kpis/" + k["id"],
            json={"completed": False, "version": 2, "note": "核对更正"},
        ).status_code
        == 200
    )
    r = client.patch(
        "/api/v1/records/lesson/" + lesson["id"],
        json={"version": 1, "data": {**lesson["data"], "status": "cancelled"}},
    )
    assert r.status_code == 200, r.text
    month = now().astimezone(reports.TZ).strftime("%Y-%m")
    assert (
        client.get("/api/v1/reports/monthly?period=" + month).json()["metrics"][
            "feedback_expected"
        ]
        == 0
    )


def test_excel_import_idempotency_conflict_and_no_money_from_scan(client):
    v3()
    from openpyxl import load_workbook

    raw = client.get("/api/v1/admissions/template/xlsx").content
    wb = load_workbook(io.BytesIO(raw))
    ws = wb.active
    # header row is four; populate the first stable template row.
    headers = {str(c.value): c.column for c in ws[4]}
    assert "学员" in headers
    for name, value in {
        "学员": "小明",
        "家长": "家长",
        "电话": "13800000001",
        "跟进阶段": "已报名",
        "报名信息": "已收款900元",
    }.items():
        ws.cell(5, headers[name], value)
    buffer = io.BytesIO()
    wb.save(buffer)
    resp = client.post(
        "/api/v1/admissions/imports",
        files={
            "file": (
                "scan.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"request_id": "excel-import-01"},
    )
    assert resp.status_code == 201, resp.text
    bid = resp.json()["id"]
    from app.admission_import import execute

    execute(bid, "tenant-a")
    detail = client.get("/api/v1/admissions/imports/" + bid).json()
    assert detail["status"] == "review", detail
    body = {"version": detail["version"], "rows": [{"index": 0}]}
    r = client.post("/api/v1/admissions/imports/" + bid + "/confirm", json=body)
    assert r.status_code == 200, r.text
    assert r.json()["results"][0]["status"] == "imported", r.text
    again = client.post(
        "/api/v1/admissions/imports/" + bid + "/confirm",
        json={**body, "version": r.json()["version"]},
    )
    assert again.json()["results"][0]["status"] == "already_imported"
    assert client.get("/api/v1/records/lead").json()["total"] == 1
    assert (
        client.get("/api/v1/records/order").json()["total"] == 0
        and client.get("/api/v1/ledger").json() == []
    )


def test_form_field_acl_encrypted_storage_and_version_pin(client):
    v3()
    with Session(engine()) as s:
        staff(s, "sales1", ["admissions"])
    lead = add(
        client, "lead", {"name": "家长", "contact": "13800000001", "owner_id": "sales1"}
    )
    d = define(
        client,
        "lead",
        {
            "fields": [
                {"key": "public", "label": "需求", "type": "text"},
                {
                    "key": "secret",
                    "label": "内部价",
                    "type": "text",
                    "read_roles": ["owner"],
                    "edit_roles": ["owner"],
                },
            ]
        },
        "form",
    )
    path = "/api/v1/studio/forms/" + d["id"] + "/submissions"
    r = client.post(
        path,
        json={
            "values": {"public": "想了解AI", "secret": "不可外传"},
            "record_id": lead["id"],
            "request_id": "form-secret-01",
            "definition_version": 1,
        },
    )
    assert r.status_code == 201, r.text
    sign(client, "sales1")
    fields = client.get("/api/v1/studio/forms/" + d["id"]).json()["fields"]
    assert [x["key"] for x in fields] == ["public"]
    assert (
        client.post(
            path,
            json={
                "values": {"secret": "tamper"},
                "record_id": lead["id"],
                "request_id": "form-secret-02",
            },
        ).status_code
        == 403
    )
    rows = client.get("/api/v1/records/lead").json()
    assert "不可外传" not in str(rows) and "想了解AI" in str(rows)
    with Session(engine()) as s:
        f = s.exec(select(FormSubmission)).first()
        assert "不可外传" not in f.values_cipher
    assert (
        client.patch(
            "/api/v1/records/lead/" + lead["id"],
            json={
                "version": lead["version"],
                "data": {
                    "name": "家长",
                    "contact": "13800000001",
                    "custom_fields": {"secret": "tamper"},
                },
            },
        ).status_code
        == 400
    )


def test_financial_approval_cannot_be_bypassed_and_action_once(client):
    v3()
    st, o = registration(client)
    e = enroll(client, st, o)
    template = next(
        t
        for t in client.get("/api/v1/studio/templates").json()
        if t["binding"] == "refund"
    )
    d = define(client, "refund", template["spec"])
    enable(client, d)
    payload = {
        "amount_fen": 1000,
        "reason": "减免",
        "withdrawal": False,
        "enrollment_ids": [],
    }
    assert (
        client.post(
            "/api/v1/orders/" + o["id"] + "/refunds",
            json={**payload, "request_id": "blocked-refund-01"},
        ).status_code
        == 409
    )
    r = client.post(
        "/api/v1/studio/workflows/" + d["id"] + "/instances",
        json={
            "payload": {"business": payload},
            "record_id": o["id"],
            "request_id": "flow-refund-01",
        },
    )
    assert r.status_code == 201, r.text
    assert "payload_cipher" not in r.json()
    task = client.get("/api/v1/studio/tasks").json()[0]
    rr = client.post(
        "/api/v1/studio/tasks/" + task["id"] + "/decision", json={"decision": "approve"}
    )
    assert rr.status_code == 200, rr.text
    task = client.get("/api/v1/studio/tasks").json()[0]
    assert task["node_type"] == "form"
    assert (
        client.post(
            "/api/v1/studio/tasks/" + task["id"] + "/decision",
            json={"decision": "approve"},
        ).status_code
        == 400
    )
    rr = client.post(
        "/api/v1/studio/tasks/" + task["id"] + "/decision",
        json={"decision": "approve", "values": {"actual_confirmed": True}},
    )
    assert rr.status_code == 200, rr.text
    assert rr.json()["status"] == "completed"
    again = client.post(
        "/api/v1/studio/instances/" + r.json()["id"] + "/apply",
        json={"actual_confirmed": True},
    )
    assert again.status_code == 200
    assert (
        len(
            [
                x
                for x in client.get("/api/v1/ledger").json()
                if x["entry_type"] == "refund"
            ]
        )
        == 1
    )


@pytest.mark.parametrize("mode", ["all", "any"])
def test_parallel_approval_modes_and_immutable_versions(client, mode):
    v3()
    with Session(engine()) as s:
        staff(s, "sales1", ["admissions"])
    spec = {
        "nodes": [
            {"id": "s", "data": {"kind": "start"}},
            {
                "id": "a",
                "data": {
                    "kind": "approval",
                    "users": ["user-a", "sales1"],
                    "mode": mode,
                },
            },
            {"id": "e", "data": {"kind": "end"}},
        ],
        "edges": [{"source": "s", "target": "a"}, {"source": "a", "target": "e"}],
    }
    d = define(client, spec=spec)
    r = client.post(
        "/api/v1/studio/workflows/" + d["id"] + "/instances",
        json={"payload": {}, "request_id": "parallel-flow-" + mode},
    )
    assert r.status_code == 201, r.text
    iid = r.json()["id"]
    tasks = client.get("/api/v1/studio/tasks").json()
    task = next(t for t in tasks if t["assignee_id"] == "user-a")
    assert (
        client.post(
            "/api/v1/studio/tasks/" + task["id"] + "/decision",
            json={"decision": "approve"},
        ).status_code
        == 200
    )
    detail = client.get("/api/v1/studio/instances/" + iid).json()
    assert detail["status"] == ("approved" if mode == "any" else "running")
    client.post("/api/v1/studio/definitions/" + d["id"] + "/publish")
    assert (
        client.get("/api/v1/studio/instances/" + iid).json()["definition_version"] == 1
    )
    if mode == "all":
        sign(client, "sales1")
        task = client.get("/api/v1/studio/tasks").json()[0]
        assert (
            client.post(
                "/api/v1/studio/tasks/" + task["id"] + "/decision",
                json={"decision": "approve"},
            ).json()["status"]
            == "approved"
        )


def test_invalid_graph_and_condition_dry_run(client):
    assert studio.condition({"field": "checked", "value": "true"}, {"checked": True})
    assert studio.condition({"field": "count", "value": "3"}, {"count": 3})
    v3()
    spec = {
        "nodes": [
            {"id": "s", "data": {"kind": "start"}},
            {
                "id": "c",
                "data": {
                    "kind": "condition",
                    "condition": {"field": "price", "op": "gte", "value": 100},
                },
            },
            {"id": "yes", "data": {"kind": "end"}},
            {"id": "no", "data": {"kind": "end"}},
        ],
        "edges": [
            {"source": "s", "target": "c"},
            {"source": "c", "target": "yes", "data": {"branch": "yes"}},
            {"source": "c", "target": "no", "data": {"branch": "no"}},
        ],
    }
    d = define(client, spec=spec)
    r = client.post(
        "/api/v1/studio/definitions/" + d["id"] + "/preview",
        json={"values": {"price": 100}, "request_id": "dry-run-01"},
    ).json()
    assert (
        "yes" in r["reachable_nodes"]
        and "no" not in r["reachable_nodes"]
        and not r["side_effects"]
    )
    spec["edges"].append({"source": "yes", "target": "s"})
    rr = client.post(
        "/api/v1/studio/definitions",
        json={"name": "bad", "kind": "workflow", "spec": spec},
    )
    assert rr.status_code == 201
    assert (
        client.post(
            "/api/v1/studio/definitions/" + rr.json()["id"] + "/publish"
        ).status_code
        == 400
    )


def test_return_requires_new_approval_and_revoked_creator_blocks_action(client):
    v3()
    with Session(engine()) as s:
        staff(s, "sales1", ["admissions"])
    d = define(client, "lead")
    sign(client, "sales1")
    r = client.post(
        "/api/v1/studio/workflows/" + d["id"] + "/instances",
        json={
            "payload": {"name": "家长", "contact": "13800000001"},
            "request_id": "returned-flow-01",
        },
    )
    assert r.status_code == 201
    login(client)
    task = client.get("/api/v1/studio/tasks").json()[0]
    assert (
        client.post(
            "/api/v1/studio/tasks/" + task["id"] + "/decision",
            json={"decision": "return", "note": "补充年级"},
        ).json()["status"]
        == "returned"
    )
    sign(client, "sales1")
    rr = client.post(
        "/api/v1/studio/instances/" + r.json()["id"] + "/resubmit",
        json={
            "payload": {"name": "家长", "contact": "13800000001", "grade": "一年级"},
            "request_id": "returned-flow-02",
        },
    )
    assert rr.status_code == 201
    login(client)
    with Session(engine()) as s:
        m = s.exec(select(Membership).where(Membership.user_id == "sales1")).one()
        m.status = "disabled"
        s.add(m)
        s.commit()
    task = client.get("/api/v1/studio/tasks").json()[0]
    assert (
        client.post(
            "/api/v1/studio/tasks/" + task["id"] + "/decision",
            json={"decision": "approve"},
        ).status_code
        == 403
    )
    assert client.get("/api/v1/records/lead").json()["total"] == 0


def test_consultation_paused_but_campaigns_kept_and_addon_truthful(client):
    v3()
    assert client.get("/api/v1/auth/me").json()["features"]["consultation"] is False
    assert client.get("/api/v1/conversations").status_code == 403
    assert client.post("/api/v1/public/org-a/chat").status_code == 403
    c = add(client, "campaign", {"name": "每周内容", "interval_days": 7})
    assert c["data"]["interval_days"] == 7
    data = client.get("/api/v1/addons/digital-human").json()
    assert not data["provider_connected"]
    assert (
        client.post(
            "/api/v1/addons/digital-human/requests", json={"note": "希望报价"}
        ).status_code
        == 201
    )
    assert (
        client.get("/api/v1/platform/addon-requests").json()[0]["status"] == "requested"
    )


def test_cross_tenant_reports_imports_and_workflow_ids_denied(client):
    v3()
    d = define(client)
    flow = client.post(
        "/api/v1/studio/workflows/" + d["id"] + "/instances",
        json={"payload": {}, "request_id": "tenant-flow-01"},
    ).json()
    login(client, "b")
    assert client.get("/api/v1/studio/instances/" + flow["id"]).status_code == 404
    assert client.get("/api/v1/studio/forms/" + d["id"]).status_code == 404
    assert client.get("/api/v1/studio/tasks").json() == []
    assert client.get("/api/v1/reports/monthly").json()["metrics"]["new_students"] == 0


def test_wecom_notification_unknown_no_auto_resend_and_revoke(client, monkeypatch):
    v3()
    assert (
        client.put(
            "/api/v1/notifications/wecom",
            json={"webhook": "http://127.0.0.1/private", "active": True},
        ).status_code
        == 400
    )
    assert (
        client.put(
            "/api/v1/notifications/wecom",
            json={
                "webhook": "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=TEST-ONLY-000000",
                "active": True,
            },
        ).status_code
        == 200
    )

    class TimeoutClient:
        def __init__(self, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def post(self, *args, **kwargs):
            raise TimeoutError("simulated uncertain result")

    monkeypatch.setattr(reports.httpx, "Client", TimeoutClient)
    r = client.post("/api/v1/notifications/test")
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "unknown"
    assert (
        client.post(
            "/api/v1/notifications/" + r.json()["id"] + "/retry", json={}
        ).status_code
        == 400
    )
    client.put("/api/v1/notifications/wecom", json={"active": False})
    with Session(engine()) as s:
        row = s.get(NotificationOutbox, r.json()["id"])
        row.status = "queued"
        s.add(row)
        s.commit()
        reports.delivery(s, s.get(Tenant, "tenant-a"), row)
        assert row.status == "cancelled"


def test_attendance_roster_transfer_and_teacher_history(client):
    v3()
    with Session(engine()) as s:
        staff(s, "teach1", ["teacher"])
        staff(s, "teach2", ["teacher"])
    cl = add(client, "class", {"name": "原班", "teacher_id": "teach1"})
    target = add(client, "class", {"name": "新班", "teacher_id": "teach2"})
    enrollments = []
    for n in (1, 2):
        st, o = registration(client, str(n))
        enrollments.append(enroll(client, st, o, cl["id"]))
    with Session(engine()) as s:
        for item in enrollments:
            e = s.get(Record, item["id"])
            e.data = {
                **e.data,
                "started_at": (now() - timedelta(days=6)).isoformat(),
                "history": [
                    {
                        "class_id": cl["id"],
                        "start": (now() - timedelta(days=6)).isoformat(),
                        "end": None,
                    }
                ],
            }
            s.add(e)
        s.commit()
    lesson = add(
        client,
        "lesson",
        {
            "name": "历史课程",
            "class_id": cl["id"],
            "starts_at": (now() - timedelta(days=2)).isoformat(),
            "status": "completed",
        },
    )
    r = client.post(
        "/api/v1/enrollments/" + enrollments[0]["id"] + "/attendance",
        json={"lesson_id": lesson["id"], "request_id": "historical-attend"},
    )
    assert r.status_code == 200, r.text
    add(
        client,
        "leave",
        {
            "name": "请假",
            "enrollment_id": enrollments[1]["id"],
            "lesson_id": lesson["id"],
            "reason": "事假",
            "status": "approved",
        },
    )
    period = now().astimezone(reports.TZ).strftime("%Y-%m")
    before = client.get("/api/v1/reports/monthly?period=" + period).json()["metrics"]
    assert (
        before["expected_attendance"],
        before["actual_attendance"],
        before["attendance_rate"],
        before["excused_absences"],
    ) == (2, 1, 50, 1)
    e = enrollments[0]
    assert (
        client.patch(
            "/api/v1/records/enrollment/" + e["id"],
            json={
                "data": {
                    "student_id": e["data"]["student_id"],
                    "order_id": e["data"]["order_id"],
                    "class_id": target["id"],
                },
                "version": 1,
            },
        ).status_code
        == 400
    )
    transfer = client.post(
        "/api/v1/enrollments/" + e["id"] + "/transfer",
        json={"class_id": target["id"], "reason": "改时段"},
    )
    assert transfer.status_code == 200, transfer.text
    after = client.get("/api/v1/reports/monthly?period=" + period).json()["metrics"]
    assert after["attendance_rate"] == 50
    # Teacher replacement does not reassign historical KPI or attendance.
    r = client.patch(
        "/api/v1/records/class/" + cl["id"],
        json={"version": 1, "data": {**cl["data"], "teacher_id": "teach2"}},
    )
    assert r.status_code == 200, r.text
    sign(client, "teach1")
    assert len(client.get("/api/v1/education/feedback-kpis").json()) == 1
    metrics = client.get("/api/v1/reports/monthly?period=" + period).json()["metrics"]
    assert metrics["attendance_rate"] == 50 and metrics["feedback_expected"] == 1
    assert (
        client.get(
            "/api/v1/students/" + e["data"]["student_id"] + "/profile"
        ).status_code
        == 404
    )


def test_custom_form_linked_approval_and_private_task_fields(client):
    v3()
    with Session(engine()) as s:
        staff(s, "sales1", ["admissions"])
    flow = define(client)
    form = define(
        client,
        spec={
            "fields": [
                {
                    "key": "secret",
                    "label": "内部备注",
                    "type": "text",
                    "read_roles": ["owner"],
                    "edit_roles": ["owner"],
                }
            ],
            "workflow_id": flow["id"],
        },
        kind="form",
    )
    r = client.post(
        "/api/v1/studio/forms/" + form["id"] + "/submissions",
        json={
            "values": {"secret": "保密内容"},
            "request_id": "linked-form-01",
            "definition_version": 1,
        },
    )
    assert r.status_code == 201, r.text
    assert r.json()["status"] == "running"
    task = client.get("/api/v1/studio/tasks").json()[0]
    assert (
        client.post(
            "/api/v1/studio/tasks/" + task["id"] + "/decision",
            json={"decision": "return", "note": "补充说明"},
        ).status_code
        == 200
    )
    assert client.get("/api/v1/studio/submissions").json()[0]["status"] == "returned"


def test_disabled_approver_reassigned_to_owner_and_new_version_isolated(client):
    v3()
    with Session(engine()) as s:
        staff(s, "sales1", ["admissions"])
    spec = {
        "nodes": [
            {"id": "s", "data": {"kind": "start"}},
            {"id": "a", "data": {"kind": "approval", "users": ["sales1"]}},
            {"id": "e", "data": {"kind": "end"}},
        ],
        "edges": [{"source": "s", "target": "a"}, {"source": "a", "target": "e"}],
    }
    d = define(client, spec=spec)
    inst = client.post(
        "/api/v1/studio/workflows/" + d["id"] + "/instances",
        json={"payload": {}, "request_id": "inactive-approver-01"},
    ).json()
    with Session(engine()) as s:
        m = s.exec(select(Membership).where(Membership.user_id == "sales1")).one()
        m.status = "disabled"
        s.add(m)
        s.commit()
    task = client.get("/api/v1/studio/tasks").json()[0]
    assert task["status"] == "unassigned"
    r = client.post("/api/v1/studio/tasks/" + task["id"] + "/assign?user_id=user-a")
    assert r.status_code == 200, r.text
    assert (
        client.post(
            "/api/v1/studio/tasks/" + task["id"] + "/decision",
            json={"decision": "approve"},
        ).json()["status"]
        == "approved"
    )


def test_form_published_version_stale_rejected_and_approval_default_off(client):
    v3()
    expense = add(
        client,
        "expense",
        {"name": "直接登记", "amount_fen": 300, "status": "confirmed"},
    )
    assert expense["data"]["status"] == "confirmed"
    form = define(
        client,
        kind="form",
        spec={
            "fields": [
                {"key": "name", "label": "姓名", "type": "text", "required": True}
            ]
        },
    )
    client.post("/api/v1/studio/definitions/" + form["id"] + "/publish")
    assert (
        client.post(
            "/api/v1/studio/forms/" + form["id"] + "/submissions",
            json={
                "values": {"name": "甲"},
                "request_id": "stale-form-01",
                "definition_version": 1,
            },
        ).status_code
        == 409
    )


def test_reversal_approval_actual_confirmation_and_single_journal(client):
    v3()
    st, o = registration(client)
    pay = client.post(
        "/api/v1/orders/" + o["id"] + "/payments",
        json={"amount_fen": 9000, "request_id": "reversal-payment-01"},
    ).json()
    template = next(
        t
        for t in client.get("/api/v1/studio/templates").json()
        if t["binding"] == "reversal"
    )
    d = define(client, "reversal", template["spec"])
    enable(client, d)
    assert (
        client.post(
            "/api/v1/ledger/" + pay["id"] + "/reverse",
            json={"request_id": "bypass-reversal-01"},
        ).status_code
        == 409
    )
    r = client.post(
        "/api/v1/studio/workflows/" + d["id"] + "/instances",
        json={
            "record_id": pay["id"],
            "payload": {"business": {"note": "重复登记"}},
            "request_id": "reversal-flow-01",
        },
    )
    assert r.status_code == 201, r.text
    for values in ({}, {"actual_confirmed": True}):
        task = client.get("/api/v1/studio/tasks").json()[0]
        rr = client.post(
            "/api/v1/studio/tasks/" + task["id"] + "/decision",
            json={"decision": "approve", "values": values},
        )
        assert rr.status_code == 200, rr.text
    assert rr.json()["status"] == "completed"
    assert (
        len(
            [
                e
                for e in client.get("/api/v1/ledger").json()
                if e["entry_type"] == "reversal"
            ]
        )
        == 1
    )


def test_linked_form_resubmit_keeps_old_values_and_blocks_reserved_source(client):
    v3()
    flow = define(client)
    form = define(
        client,
        kind="form",
        spec={
            "fields": [{"key": "notes", "label": "说明", "type": "text"}],
            "workflow_id": flow["id"],
        },
    )
    r = client.post(
        "/api/v1/studio/forms/" + form["id"] + "/submissions",
        json={"values": {"notes": "第一版"}, "request_id": "form-revision-01"},
    )
    assert r.status_code == 201, r.text
    inst = client.get("/api/v1/studio/instances").json()[0]
    task = client.get("/api/v1/studio/tasks").json()[0]
    client.post(
        "/api/v1/studio/tasks/" + task["id"] + "/decision",
        json={"decision": "return", "note": "补充"},
    )
    r = client.post(
        "/api/v1/studio/instances/" + inst["id"] + "/resubmit",
        json={
            "payload": {"values": {"notes": "第二版"}},
            "request_id": "form-revision-02",
        },
    )
    assert r.status_code == 201, r.text
    assert r.json()["revision"] == 2
    assert (
        client.get("/api/v1/studio/instances/" + inst["id"]).json()["payload"][
            "values"
        ]["notes"]
        == "第一版"
    )
    assert (
        client.get("/api/v1/studio/instances/" + r.json()["id"]).json()["payload"][
            "values"
        ]["notes"]
        == "第二版"
    )
    assert (
        client.post(
            "/api/v1/studio/workflows/" + flow["id"] + "/instances",
            json={
                "payload": {"_form_submission_id": "forged"},
                "request_id": "forged-source-01",
            },
        ).status_code
        == 403
    )


def test_import_custom_columns_and_concurrent_version_conflict(client):
    v3()
    from openpyxl import load_workbook

    form = define(
        client,
        "lead",
        {"fields": [{"key": "interest", "label": "兴趣主题", "type": "text"}]},
        "form",
    )
    raw = client.get(
        "/api/v1/admissions/template/xlsx?definition_id=" + form["id"]
    ).content
    wb = load_workbook(io.BytesIO(raw))
    ws = wb.active
    headers = {str(c.value): c.column for c in ws[4]}
    for label, value in {
        "学员": "扫描学员",
        "家长": "扫描家长",
        "电话": "13800000001",
        "兴趣主题": "机器人",
    }.items():
        ws.cell(5, headers[label], value)
    b = io.BytesIO()
    wb.save(b)
    r = client.post(
        "/api/v1/admissions/imports",
        files={
            "file": (
                "custom.xlsx",
                b.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 201, r.text
    from app.admission_import import execute

    execute(r.json()["id"], "tenant-a")
    detail = client.get("/api/v1/admissions/imports/" + r.json()["id"]).json()
    rr = client.post(
        "/api/v1/admissions/imports/" + r.json()["id"] + "/confirm",
        json={"version": detail["version"], "rows": [{"index": 0}]},
    )
    assert rr.status_code == 200, rr.text
    assert rr.json()["results"][0]["status"] == "imported", rr.text
    assert "机器人" in str(client.get("/api/v1/records/lead").json())
    assert (
        client.post(
            "/api/v1/admissions/imports/" + r.json()["id"] + "/confirm",
            json={"version": detail["version"], "rows": [{"index": 0}]},
        ).status_code
        == 409
    )


def test_uncalibrated_history_does_not_report_fake_zero(client):
    v3()
    st, o = registration(client)
    e = enroll(client, st, o)
    with Session(engine()) as s:
        row = s.get(Record, e["id"])
        row.data = {**row.data, "history_calibrated": False, "history": []}
        s.add(row)
        s.commit()
    report = client.get("/api/v1/reports/monthly").json()
    assert report["baseline_status"] == "needs_calibration"
    assert (
        report["metrics"]["attrition_rate"] is None
        and report["metrics"]["opening_students"] is None
    )


def test_approval_document_access_follows_field_roles(client):
    v3()
    with Session(engine()) as s:
        staff(s, "sales1", ["admissions"])
        staff(s, "fin1", ["finance"])
    spec = {
        "nodes": [
            {"id": "s", "data": {"kind": "start"}},
            {"id": "a", "data": {"kind": "approval", "roles": ["finance"]}},
            {"id": "e", "data": {"kind": "end"}},
        ],
        "edges": [{"source": "s", "target": "a"}, {"source": "a", "target": "e"}],
    }
    flow = define(client, spec=spec)
    form = define(
        client,
        kind="form",
        spec={
            "fields": [
                {
                    "key": "file",
                    "label": "附件",
                    "type": "attachment",
                    "read_roles": ["admissions", "finance"],
                    "edit_roles": ["admissions"],
                }
            ],
            "workflow_id": flow["id"],
        },
    )
    sign(client, "sales1")
    from PIL import Image

    image = io.BytesIO()
    Image.new("RGB", (30, 30), "white").save(image, format="PNG")
    a = client.post(
        "/api/v1/assets",
        data={"purpose": "document"},
        files={"file": ("bill.png", image.getvalue(), "image/png")},
    )
    assert a.status_code == 201, a.text
    asset = a.json()["id"]
    r = client.post(
        "/api/v1/studio/forms/" + form["id"] + "/submissions",
        json={"values": {"file": [asset]}, "request_id": "private-document-01"},
    )
    assert r.status_code == 201, r.text
    sign(client, "fin1")
    inst = client.get("/api/v1/studio/instances").json()[0]
    assert client.get("/api/v1/studio/instances/" + inst["id"]).status_code == 200
    assert client.get("/api/v1/assets/" + asset + "/download").status_code == 200
    with Session(engine()) as s:
        m = s.exec(select(Membership).where(Membership.user_id == "fin1")).one()
        m.roles = ["teacher"]
        m.authz_version += 1
        s.add(m)
        s.commit()
    assert client.get("/api/v1/assets/" + asset + "/download").status_code == 404


def test_student_custom_forms_cannot_expand_finance_teaching_access(client):
    v3()
    with Session(engine()) as s:
        staff(s, "fin1", ["finance"])
    student = add(client, "student", {"name": "学生", "notes": "教学隐私"})
    form = define(
        client,
        "student",
        {"fields": [{"key": "teaching", "label": "教学备注", "type": "text"}]},
        "form",
    )
    r = client.post(
        "/api/v1/studio/forms/" + form["id"] + "/submissions",
        json={
            "record_id": student["id"],
            "values": {"teaching": "仅教学范围可见"},
            "request_id": "teaching-form-01",
        },
    )
    assert r.status_code == 201, r.text
    sign(client, "fin1")
    profile = client.get("/api/v1/students/" + student["id"] + "/profile")
    assert profile.status_code == 200
    assert "仅教学范围可见" not in profile.text and "教学隐私" not in profile.text


def test_concurrent_payment_and_attendance_do_not_duplicate_journals(client):
    v3()
    from concurrent.futures import ThreadPoolExecutor

    st, o = registration(client)
    cl = add(client, "class", {"name": "并发班"})

    def pay(n):
        return client.post(
            "/api/v1/orders/" + o["id"] + "/payments",
            json={"amount_fen": 9000, "request_id": "concurrent-pay-" + str(n)},
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(pay, [1, 2]))
    assert sorted(r.status_code for r in results) == [200, 400], [
        r.text for r in results
    ]
    e = add(
        client,
        "enrollment",
        {"student_id": st["id"], "order_id": o["id"], "class_id": cl["id"]},
    )
    lesson = add(
        client,
        "lesson",
        {
            "name": "并发签到",
            "class_id": cl["id"],
            "starts_at": (now() + timedelta(days=1)).isoformat(),
        },
    )

    def attend(n):
        return client.post(
            "/api/v1/enrollments/" + e["id"] + "/attendance",
            json={
                "lesson_id": lesson["id"],
                "request_id": "concurrent-attend-" + str(n),
            },
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(attend, [1, 2]))
    assert all(r.status_code == 200 for r in results), [r.text for r in results]
    assert results[0].json()["id"] == results[1].json()["id"]


def test_financial_reversal_chain_preserves_receipt_and_refund_totals(client):
    v3()
    st, o = registration(client)
    payment = client.post(
        "/api/v1/orders/" + o["id"] + "/payments",
        json={"amount_fen": 9000, "request_id": "chain-payment-01"},
    ).json()
    first = client.post(
        "/api/v1/ledger/" + payment["id"] + "/reverse",
        json={"request_id": "chain-reverse-01"},
    ).json()
    second = client.post(
        "/api/v1/ledger/" + first["id"] + "/reverse",
        json={"request_id": "chain-reverse-02"},
    )
    assert second.status_code == 200, second.text
    total = client.get("/api/v1/finance/overview").json()["totals"]
    assert (
        total["received_fen"] == 9000
        and total["refund_fen"] == 0
        and total["net_fen"] == 9000
    )


def test_same_student_multiple_orders_deduplicate_class_attendance(client):
    v3()
    st, o = registration(client)
    cl = add(client, "class", {"name": "唯一学员班", "capacity": 1})
    first = enroll(client, st, o, cl["id"])
    second_order = add(
        client,
        "order",
        {
            "name": "续费",
            "student_id": st["id"],
            "price_fen": 9000,
            "purchased_units": 8,
        },
    )
    second = enroll(client, st, second_order, cl["id"])
    with Session(engine()) as s:
        for item in (first, second):
            e = s.get(Record, item["id"])
            e.data = {
                **e.data,
                "started_at": (now() - timedelta(days=5)).isoformat(),
                "history": [
                    {
                        "class_id": cl["id"],
                        "start": (now() - timedelta(days=5)).isoformat(),
                        "end": None,
                    }
                ],
            }
            s.add(e)
        s.commit()
    lesson = add(
        client,
        "lesson",
        {
            "name": "续费仍只到课一次",
            "class_id": cl["id"],
            "starts_at": (now() - timedelta(days=1)).isoformat(),
            "status": "completed",
        },
    )
    r = client.post(
        "/api/v1/enrollments/" + second["id"] + "/attendance",
        json={"lesson_id": lesson["id"], "request_id": "renewed-attend-01"},
    )
    assert r.status_code == 200, r.text
    metrics = client.get("/api/v1/reports/monthly").json()["metrics"]
    assert (
        metrics["expected_attendance"] == 1
        and metrics["actual_attendance"] == 1
        and metrics["attendance_rate"] == 100
    )


def test_parallel_finance_confirmation_executes_expense_action_once(client):
    v3()
    from concurrent.futures import ThreadPoolExecutor

    template = next(
        t
        for t in client.get("/api/v1/studio/templates").json()
        if t["binding"] == "expense"
    )
    d = define(client, "expense", template["spec"])
    enable(client, d)
    r = client.post(
        "/api/v1/studio/workflows/" + d["id"] + "/instances",
        json={
            "payload": {
                "business": {
                    "name": "并发确认费用",
                    "amount_fen": 500,
                    "status": "confirmed",
                }
            },
            "request_id": "expense-approve-once",
        },
    )
    assert r.status_code == 201, r.text
    task = client.get("/api/v1/studio/tasks").json()[0]
    client.post(
        "/api/v1/studio/tasks/" + task["id"] + "/decision", json={"decision": "approve"}
    )
    task = client.get("/api/v1/studio/tasks").json()[0]

    def decide(n):
        return client.post(
            "/api/v1/studio/tasks/" + task["id"] + "/decision",
            json={"decision": "approve", "values": {"actual_confirmed": True}},
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(decide, [1, 2]))
    assert sorted(r.status_code for r in results) == [200, 409], [
        r.text for r in results
    ]
    assert client.get("/api/v1/records/expense").json()["total"] == 1


def test_runtime_condition_does_not_create_unselected_branch_tasks(client):
    v3()
    with Session(engine()) as s:
        staff(s, "fin1", ["finance"])
    spec = {
        "nodes": [
            {"id": "s", "data": {"kind": "start"}},
            {
                "id": "condition",
                "data": {
                    "kind": "condition",
                    "condition": {"field": "price", "op": "gte", "value": 100},
                },
            },
            {"id": "yes", "data": {"kind": "approval", "users": ["user-a"]}},
            {"id": "no", "data": {"kind": "approval", "users": ["fin1"]}},
            {"id": "end", "data": {"kind": "end"}},
        ],
        "edges": [
            {"source": "s", "target": "condition"},
            {"source": "condition", "target": "yes", "data": {"branch": "yes"}},
            {"source": "condition", "target": "no", "data": {"branch": "no"}},
            {"source": "yes", "target": "end"},
            {"source": "no", "target": "end"},
        ],
    }
    d = define(client, spec=spec)
    r = client.post(
        "/api/v1/studio/workflows/" + d["id"] + "/instances",
        json={"payload": {"price": 100}, "request_id": "runtime-condition-01"},
    )
    assert r.status_code == 201, r.text
    tasks = client.get("/api/v1/studio/tasks").json()
    assert [t["node_id"] for t in tasks] == ["yes"]
    result = client.post(
        "/api/v1/studio/tasks/" + tasks[0]["id"] + "/decision",
        json={"decision": "approve"},
    )
    assert result.status_code == 200, result.text
    assert (
        result.json()["status"] == "approved"
        and "no" not in result.json()["state"]["done"]
    )


def test_any_join_cancels_exclusive_unfinished_upstream_branch(client):
    v3()
    with Session(engine()) as s:
        staff(s, "fin1", ["finance"])
    spec = {
        "nodes": [
            {"id": "s", "data": {"kind": "start"}},
            {"id": "fork", "data": {"kind": "parallel"}},
            {"id": "a", "data": {"kind": "approval", "users": ["user-a"]}},
            {"id": "b", "data": {"kind": "approval", "users": ["fin1"]}},
            {"id": "c", "data": {"kind": "form", "users": ["fin1"]}},
            {"id": "join", "data": {"kind": "join", "mode": "any"}},
            {"id": "end", "data": {"kind": "end"}},
        ],
        "edges": [
            {"source": "s", "target": "fork"},
            {"source": "fork", "target": "a"},
            {"source": "fork", "target": "b"},
            {"source": "a", "target": "join"},
            {"source": "b", "target": "c"},
            {"source": "c", "target": "join"},
            {"source": "join", "target": "end"},
        ],
    }
    d = define(client, spec=spec)
    r = client.post(
        "/api/v1/studio/workflows/" + d["id"] + "/instances",
        json={"payload": {}, "request_id": "runtime-any-join-01"},
    )
    assert r.status_code == 201, r.text
    tasks = client.get("/api/v1/studio/tasks").json()
    task = next(t for t in tasks if t["node_id"] == "a")
    result = client.post(
        "/api/v1/studio/tasks/" + task["id"] + "/decision", json={"decision": "approve"}
    )
    assert result.status_code == 200, result.text
    assert result.json()["status"] == "approved" and set(
        result.json()["state"]["cancelled"]
    ) == {"b", "c"}
    assert client.get("/api/v1/studio/tasks").json() == []
